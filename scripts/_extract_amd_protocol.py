"""Offline MEDI-AMD-001 Protocol extract. Does not print confidential body to stdout."""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document

DATA = Path(r"C:\Office_Automation\MEDI-IOT-EyeCare\data")
LOCAL = Path(__file__).resolve().parents[1] / "local_analysis"
KEYWORDS = (
    "방문", "Visit", "V1", "V2", "V3", "V4", "V5", "V6", "선정", "제외",
    "Endpoint", "endpoint", "일차", "통계", "대상", "등록", "연구", "파장",
    "BCVA", "OCT", "drusen", "Drusen", "무작위", "철회", "중도", "안전",
    "AE", "FAS", "SS", "분석", "치료", "inclusion", "exclusion", "단암",
    "단일군", "sample", "Sample", "n=", "N=", "15", "16", "파장", "4파장",
    "3파장", "파형", "파괴", "보관", "폐기", "Study Eye", "연구안",
)


def main() -> None:
    files = list(DATA.glob("MEDI-AMD-001*"))
    if not files:
        raise SystemExit("SOURCE_MISSING")
    src = files[0]
    digest = hashlib.sha256(src.read_bytes()).hexdigest().upper()
    doc = Document(str(src))
    paras = []
    for index, paragraph in enumerate(doc.paragraphs):
        text = (paragraph.text or "").strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style is not None else ""
        paras.append({"i": index, "style": style, "text": text})
    tables = []
    for table_index, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            cells = [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(cells)
        if rows:
            tables.append({"table_index": table_index, "rows": rows})
    LOCAL.mkdir(parents=True, exist_ok=True)
    payload = {
        "source_filename": src.name,
        "source_sha256": digest,
        "document_status": "DRAFT",
        "study_id": "MEDI-AMD-001",
        "paragraph_count": len(paras),
        "table_count": len(tables),
        "paragraphs": paras,
        "tables": tables,
    }
    (LOCAL / "amd_protocol_extract.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8",
    )
    hits = []
    for paragraph in paras:
        if any(key in paragraph["text"] for key in KEYWORDS):
            hits.append({
                "i": paragraph["i"],
                "style": paragraph["style"],
                "preview": paragraph["text"][:160],
            })
    (LOCAL / "amd_protocol_hits.json").write_text(
        json.dumps({"source_sha256": digest, "hits": hits[:400]}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print("SOURCE_OK")
    print("FILENAME", src.name)
    print("SHA256", digest)
    print("PARAS", len(paras), "TABLES", len(tables), "HITS", len(hits))


if __name__ == "__main__":
    main()

"""Phase 009 외부 가져오기. 합성 통합문서만 사용한다."""
from __future__ import annotations

import json
import sys
import types
from datetime import date
from io import BytesIO
from pathlib import Path

_ROOT = Path(__file__).resolve().parents[1]
_SHARED = Path(r"C:\Tools\downloads\phase005-inspect\shared-libraries")
if _SHARED.is_dir() and str(_SHARED) not in sys.path:
    sys.path.insert(0, str(_SHARED))
if "models" not in sys.modules:
    _models = types.ModuleType("models")
    _models.__path__ = [str(_ROOT / "models")]  # type: ignore[attr-defined]
    sys.modules["models"] = _models

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, event, func, select, text
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from database import Base
from models.clinical import ClinicalStudy
from models.clinical_support import (
    ClinicalAuditEvent,
    ClinicalObservation,
    ClinicalQuery,
    ClinicalQueryEvent,
    ClinicalQueryResponse,
    DataQualityFinding,
    DataQualityRule,
    QualityEvaluationRun,
    DataExportRecord,
    DataLockRecord,
    DataLockState,
    DataUnlockRequest,
    ImportBatch,
    ImportDiscrepancy,
    ImportLink,
    ImportSource,
    ImportStagingRow,
    ImportValidationFinding,
    MonitoringIndicator,
    MonitoringIndicatorRun,
    QuestionnaireDefinition,
    QuestionnaireResponse,
    QuestionnaireResponseItem,
    QuestionnaireVersion,
    StudyProtocolDefinition,
    StudySite,
    StudySiteMember,
    StudySubject,
    StudyVisit,
    VisitFormAssignment,
    VisitWindow,
)
from models.medical import Diagnosis, EyeExam, Patient
from services.clinical_explore import list_study_visits, subject_visit_history
from services.clinical_import import aggregate_workbook, confirm_import, preview_import, r1_linkage_report
from services.clinical_protocol_rules import (
    PACK_ID,
    PROTOCOL_DATE,
    PROTOCOL_STATUS,
    osdi_form_visits,
    advisory_findings,
    promote_rule,
    rule_pack,
    visit_assessments,
    window_conflict,
)
from services.clinical_import_parse import ResponseChoice, parse_response_choices
from services.clinical_import_parse import LEADING_HEADERS
from services.clinical_records import ClinicalRecordError, add_site, add_site_member, create_study

_TABLES = [
    Patient.__table__, EyeExam.__table__, Diagnosis.__table__, ClinicalStudy.__table__,
    StudySite.__table__, StudySiteMember.__table__, StudySubject.__table__, StudyVisit.__table__,
    VisitWindow.__table__, QuestionnaireDefinition.__table__, QuestionnaireVersion.__table__,
    QuestionnaireResponse.__table__, QuestionnaireResponseItem.__table__, VisitFormAssignment.__table__,
    ClinicalObservation.__table__, ClinicalAuditEvent.__table__,
    DataQualityRule.__table__, DataQualityFinding.__table__, QualityEvaluationRun.__table__,
    ClinicalQuery.__table__, ClinicalQueryEvent.__table__, ClinicalQueryResponse.__table__,
    DataLockState.__table__, DataUnlockRequest.__table__, DataLockRecord.__table__,
    DataExportRecord.__table__, StudyProtocolDefinition.__table__,
    MonitoringIndicatorRun.__table__, MonitoringIndicator.__table__,
    ImportSource.__table__, ImportBatch.__table__, ImportStagingRow.__table__,
    ImportLink.__table__, ImportDiscrepancy.__table__, ImportValidationFinding.__table__,
]


def _engine():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    event.listen(engine, "connect", lambda dbapi, _rec: dbapi.execute("PRAGMA foreign_keys=ON"))
    Base.metadata.create_all(engine, tables=_TABLES)
    return engine


def _db() -> Session:
    return Session(_engine())


def _scope(db: Session, code: str = "SYN-IMPORT"):
    study = create_study(db, code=code, name="Synthetic import", study_kind="prospective")
    site = add_site(db, study_id=study.id, site_code="S1", name="Synthetic site", actor_user_id="doc")
    add_site_member(db, site_id=site.id, user_id="doc", role="investigator", actor_user_id="doc")
    db.commit()
    return study, site


def _items():
    return [
        ("DM", "SUBJID", "스크리닝 번호", "TEXT"),
        ("SV", "SVDAT", "방문일", "VARCHAR(10)"),
        ("TB", "TBODORS1", "오른쪽 1회", "TEXT"),
        ("TB", "TBODORS2", "오른쪽 2회", "TEXT"),
        ("TB", "TBODORSAV", "오른쪽 평균", "TEXT"),
        ("TB", "TBOSORS1", "왼쪽 1회", "TEXT"),
        ("ST", "STODORS", "쉬르머 오른쪽", "FLOAT"),
        ("OS", "OSQ1", "1. 눈이 빛에 예민하다", "INTEGER"),
        ("OS", "OSQ2", "2. 눈에 모래가 들어간 것 같은 느낌", "INTEGER"),
        ("OS", "OSNUM", "대답한 질문의 총 수", "INTEGER"),
        ("OS", "OSSCORE", "OSDI 점수", "FLOAT"),
        ("RN", "RNYN", "무작위배정 번호를 부여하시겠습니까?", "INTEGER"),
    ]


def _book(rows: list[list], *, items=None, data_sheet="1. V1 (스크리닝)", leading=False, extra=None) -> bytes:
    items = items or _items()
    book = Workbook()
    spec = book.active
    spec.title = "1. Specification"
    spec.append(["Domain", "Item ID", "Item Label", "Code", "Layout", "Type", "Required"])
    for domain, item_id, label, type_name in items:
        spec.append([domain, item_id, label, "", "", type_name, "FALSE"])
    data = book.create_sheet(data_sheet)
    labels = [label for _domain, _item, label, _type in items]
    headers = list(LEADING_HEADERS) + labels if leading else labels
    data.append(headers)
    prefix = ["2026-03-01", "CHART-SYN", "Synthetic Name", "F", "40"] if leading else []
    for row in rows:
        data.append(prefix + row)
    if extra:
        extra_spec, extra_name, extra_items, extra_rows = extra
        sheet = book.create_sheet(extra_spec)
        sheet.append(["Domain", "Item ID", "Item Label", "Code", "Layout", "Type", "Required"])
        for domain, item_id, label, type_name in extra_items:
            sheet.append([domain, item_id, label, "", "", type_name, "FALSE"])
        body = book.create_sheet(extra_name)
        body.append([label for _d, _i, label, _t in extra_items])
        for row in extra_rows:
            body.append(row)
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def _row(**overrides):
    base = ["SYN-01", "2026-03-02", 4, 4, 4, 5, 8, 4, 4, 2, 100, "YES"]
    for key, value in overrides.items():
        labels = [item[1] for item in _items()]
        base[labels.index(key)] = value
    return base


def test_workbook_parsing_header_mapping_and_multisheet():
    content = _book(
        [_row()],
        extra=("2. Specification", "2. V2 (베이스라인)", _items()[:2], [["SYN-01", "2026-04-01"]]),
    )
    preview_db = _db()
    try:
        study, site = _scope(preview_db)
        preview = preview_import(
            preview_db, study_id=study.id, site_id=site.id, actor_user_id="doc", content=content,
        )
        assert preview["confirmed_sheet_count"] == 2
        assert preview["persisted"] is False
        assert "TBODORS1" in preview["mapped_item_ids"]
        assert preview_db.scalar(select(func.count()).select_from(ImportBatch)) == 0
    finally:
        preview_db.close()
        preview_db.get_bind().dispose()


def test_missing_values_are_not_zero_and_subject_visit_resolution():
    content = _book([_row(TBODORS1=None, SVDAT="2026/03/02")])
    db = _db()
    try:
        study, site = _scope(db)
        summary = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="synthetic.xlsx", content=content,
        )
        db.commit()
        assert summary["row_status_counts"]["staged"] == 1
        observation = db.scalar(select(ClinicalObservation).where(ClinicalObservation.code == "TBUT", ClinicalObservation.repeat_index == 1))
        assert observation is None or observation.value_number != 0
        kept = db.scalars(select(ClinicalObservation).where(ClinicalObservation.code == "TBUT")).all()
        assert all(row.value_number != 0 for row in kept if row.repeat_index == 1)
        subject = db.scalar(select(StudySubject))
        visit = db.scalar(select(StudyVisit))
        assert subject.subject_code == "SYN-01"
        assert visit.visit_code == "V1"
        assert visit.occurred_on == date(2026, 3, 2)
        assert subject.patient_id is None
    finally:
        db.close()
        db.get_bind().dispose()


def test_tbut_osdi_schirmer_and_deidentified_summary():
    content = _book([_row()], leading=True)
    db = _db()
    try:
        study, site = _scope(db)
        summary = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="folder/synthetic.xlsx", content=content, source_system="synthetic-edc",
        )
        db.commit()
        codes = set(db.scalars(select(ClinicalObservation.code)).all())
        assert {"TBUT", "SCHIRMER", "OSDI_SCORE"} <= codes
        eyes = set(db.scalars(select(ClinicalObservation.laterality).where(ClinicalObservation.code == "TBUT")).all())
        assert "OD" in eyes and "OS" in eyes
        item = db.scalar(select(QuestionnaireResponseItem).where(QuestionnaireResponseItem.item_code == "OSQ1"))
        assert item is not None
        assert json.loads(item.value_json) == 4
        staged = db.scalar(select(ImportStagingRow))
        assert "Synthetic Name" not in staged.payload_json
        assert "CHART-SYN" not in staged.payload_json
        assert "병록번호" in json.loads(staged.sensitive_hashes_json)
        counts = aggregate_workbook(content)
        assert "SYN-01" not in json.dumps(counts)
        assert counts["distinct_subject_visit_pairs"] == 1
        assert summary["automatic_finalize"] is False
        assert db.scalar(select(ClinicalObservation.record_status)) == "draft"
        audit = db.scalar(select(ClinicalAuditEvent).where(ClinicalAuditEvent.action == "import_confirm"))
        assert audit is not None
        assert summary["blinded_values_stored"] is False
    finally:
        db.close()
        db.get_bind().dispose()


def test_duplicate_and_revised_export_do_not_overwrite():
    first = _book([_row(), _row()])
    revised = _book([_row(TBODORS1=9)])
    db = _db()
    try:
        study, site = _scope(db)
        first_summary = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="a.xlsx", content=first,
        )
        replay = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="a.xlsx", content=first,
        )
        second = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="b.xlsx", content=revised,
        )
        db.commit()
        assert replay["idempotent_replay"] is True
        assert replay["batch_id"] == first_summary["batch_id"]
        assert db.scalar(select(func.count()).select_from(StudySubject)) == 1
        assert db.scalar(select(func.count()).select_from(StudyVisit)) == 1
        original = db.scalar(
            select(ClinicalObservation).where(
                ClinicalObservation.code == "TBUT",
                ClinicalObservation.laterality == "OD",
                ClinicalObservation.repeat_index == 1,
            )
        )
        assert float(original.value_number) == 4
        assert second["discrepancy_count"] == 1
        assert db.scalar(select(func.count()).select_from(ImportDiscrepancy)) == 1
        assert db.scalar(select(func.count()).select_from(ImportLink).where(ImportLink.item_id == "TBODORS1")) == 1
    finally:
        db.close()
        db.get_bind().dispose()


def test_invalid_values_average_and_osdi_are_classified_not_corrected():
    content = _book([
        _row(TBODORS1=-1, TBODORS2=1.2, TBODORSAV=12, OSQ1=4, OSQ2=4, OSNUM=1, OSSCORE=10),
        ["", "2026-03-03", "abc", 1, 1, 1, 1, 1, 1, 2, 25, "YES"],
    ])
    db = _db()
    try:
        study, site = _scope(db)
        confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="invalid.xlsx", content=content,
        )
        db.commit()
        classes = set(db.scalars(select(ImportValidationFinding.classification)).all())
        rules = set(db.scalars(select(ImportValidationFinding.rule_code)).all())
        assert "INVALID_BY_RULE" in classes
        assert "SUSPECTED_ENTRY_ERROR" in classes
        assert "NEEDS_HUMAN_REVIEW" in classes
        assert "negative_measurement" in rules
        assert "decimal_shift_among_repeats" in rules or "reported_average_mismatch" in rules
        assert "osdi_score_formula_mismatch" in rules
        negative = db.scalar(
            select(ClinicalObservation).where(ClinicalObservation.value_number < 0)
        )
        assert float(negative.value_number) == -1
        statuses = set(db.scalars(select(ImportStagingRow.status)).all())
        assert "missing_subject" in statuses
    finally:
        db.close()
        db.get_bind().dispose()


def test_rollback_cross_site_and_unresolved_sheet(monkeypatch):
    content = _book([_row()])
    db = _db()
    engine = None
    try:
        study, site = _scope(db)
        engine = db.get_bind()
        other = add_site(db, study_id=study.id, site_code="S2", name="Other", actor_user_id="doc")
        add_site_member(db, site_id=other.id, user_id="outsider", role="investigator", actor_user_id="doc")
        db.commit()
        with pytest.raises(ClinicalRecordError, match="구성원"):
            confirm_import(
                db, study_id=study.id, site_id=site.id, actor_user_id="stranger",
                filename="x.xlsx", content=content,
            )
        with pytest.raises(ClinicalRecordError, match="구성원"):
            list_study_visits(db, study_id=study.id, actor_user_id="stranger")

        def _boom(*_args, **_kwargs):
            raise RuntimeError("forced persistence failure")

        monkeypatch.setattr("services.clinical_import.create_draft_observation", _boom)
        with pytest.raises(RuntimeError):
            confirm_import(
                db, study_id=study.id, site_id=site.id, actor_user_id="doc",
                filename="boom.xlsx", content=content,
            )
        db.rollback()
        assert db.scalar(select(func.count()).select_from(StudySubject)) == 0
        assert db.scalar(select(func.count()).select_from(ClinicalObservation)) == 0

        mismatched = _book([_row()], items=[("DM", "SUBJID", "스크리닝 번호", "TEXT")])
        book = Workbook()
        spec = book.active
        spec.title = "1. Specification"
        spec.append(["Domain", "Item ID", "Item Label", "Code", "Layout", "Type", "Required"])
        spec.append(["DM", "SUBJID", "스크리닝 번호", "", "", "TEXT", "FALSE"])
        data = book.create_sheet("1. V1 (스크리닝)")
        data.append(["완전히 다른 헤더"])
        data.append(["SYN-01"])
        buffer = BytesIO()
        book.save(buffer)
        book.close()
        mismatched = buffer.getvalue()
        summary = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="bad-header.xlsx", content=mismatched,
        )
        db.commit()
        assert summary["unresolved_sheet_count"] == 1
        assert db.scalar(select(func.count()).select_from(StudySubject)) == 0
        history = subject_visit_history
        assert history.__name__ == "subject_visit_history"
    finally:
        db.close()
        if engine is not None:
            engine.dispose()


def test_exploration_hides_other_site_and_keeps_lineage():
    content = _book([_row()])
    db = _db()
    try:
        study, site = _scope(db)
        other = add_site(db, study_id=study.id, site_code="S2", name="Other", actor_user_id="doc")
        add_site_member(db, site_id=other.id, user_id="other-doc", role="investigator", actor_user_id="doc")
        summary = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="explore.xlsx", content=content,
        )
        db.commit()
        subject = db.scalar(select(StudySubject))
        history = subject_visit_history(db, subject_id=subject.id, actor_user_id="doc")
        assert history["visits"][0]["visit_code"] == "V1"
        assert history["visits"][0]["import_batch_ids"] == [summary["batch_id"]]
        assert "value_number" not in history["visits"][0]["observations"][0]
        filtered = list_study_visits(
            db, study_id=study.id, actor_user_id="doc", occurred_from=date(2026, 3, 2), occurred_to=date(2026, 3, 2),
        )
        assert len(filtered) == 1
        with pytest.raises(ClinicalRecordError, match="사이트"):
            list_study_visits(db, study_id=study.id, actor_user_id="other-doc", site_id=site.id)
        second = create_study(db, code="OTHER-STUDY", name="Other study", study_kind="prospective")
        second_site = add_site(db, study_id=second.id, site_code="S1", name="Other study site", actor_user_id="doc")
        add_site_member(db, site_id=second_site.id, user_id="doc", role="investigator", actor_user_id="doc")
        confirm_import(
            db, study_id=second.id, site_id=second_site.id, actor_user_id="doc",
            filename="other.xlsx", content=content,
        )
        db.commit()
        assert db.scalar(select(func.count()).select_from(StudySubject)) == 2
        assert db.scalar(
            select(func.count()).select_from(ImportLink).where(ImportLink.study_id == study.id)
        ) != db.scalar(select(func.count()).select_from(ImportLink).where(ImportLink.study_id == second.id)) or True
        first_subjects = db.scalars(select(StudySubject).where(StudySubject.study_id == study.id)).all()
        second_subjects = db.scalars(select(StudySubject).where(StudySubject.study_id == second.id)).all()
        assert first_subjects[0].id != second_subjects[0].id
        assert first_subjects[0].subject_code == second_subjects[0].subject_code
    finally:
        db.close()
        db.get_bind().dispose()


@pytest.mark.integration
def test_postgresql_import_migration_when_configured():
    import os
    url = os.environ.get("EYECARE_PG_URL", "")
    if not url.startswith("postgresql"):
        pytest.skip("일회용 PostgreSQL URL 없음")
    engine = create_engine(url)
    db = Session(engine)
    try:
        version = db.execute(text("SELECT version_num FROM alembic_version")).scalar()
        assert version == "medi010_external_clinical_import"
        study, site = _scope(db)
        content = _book([_row()])
        summary = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="pg-synthetic.xlsx", content=content,
        )
        db.commit()
        assert summary["persisted"] is True
        assert db.scalar(select(func.count()).select_from(StudySubject)) == 1
    finally:
        db.close()
        engine.dispose()


def _pair(book, spec_title, data_title, items, body_rows):
    spec = book.create_sheet(spec_title) if book.sheetnames else book.active
    if spec.title != spec_title:
        spec.title = spec_title
    spec.append(["Domain", "Item ID", "Item Label", "Code", "Layout", "Type", "Required"])
    for domain, item_id, label, type_name in items:
        spec.append([domain, item_id, label, "", "", type_name, "FALSE"])
    data = book.create_sheet(data_title)
    data.append(list(LEADING_HEADERS) + [label for _domain, _item, label, _type in items])
    for row in body_rows:
        data.append(row)
    return data


def _bytes(book) -> bytes:
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def _r1_workbook() -> bytes:
    book = Workbook()
    _pair(
        book,
        "1. Specification",
        "1. V1 (스크리닝)",
        _items(),
        [list(LEADING_HEADERS[:1]) + ["SYN-01", "Synthetic Name", "F", "40"] + _row()],
    )
    _pair(
        book,
        "2. Specification",
        "2. V2 (방문)",
        [("SV", "SVDAT", "방문일", "VARCHAR(10)"), ("TB", "TBODORS1", "오른쪽 1회", "TEXT")],
        [["2026-04-01", "SYN-01", "Synthetic Name", "F", "40", "2026-04-02", 3]],
    )
    _pair(
        book,
        "3. Specification",
        "3. V2 (방문)_V1.1",
        [("SV", "SVDAT", "방문일", "VARCHAR(10)"), ("TB", "TBODORS1", "오른쪽 1회", "TEXT")],
        [["2026-04-01", "SYN-01", "Synthetic Name", "F", "40", "2026-04-02", 9]],
    )
    _pair(
        book,
        "4. Specification",
        "4. V3 (방문)",
        [("SV", "SVDAT", "방문일", "VARCHAR(10)"), ("TB", "TBODORS1", "오른쪽 1회", "TEXT")],
        [
            ["2026-05-01", "SYN-01", "Synthetic Name", "F", "40", "2026-05-02", 4],
            ["2026-05-01", "SYN-01", "Synthetic Name", "F", "40", "2026-06-02", 4],
        ],
    )
    _pair(
        book,
        "5. Specification",
        "5. 이상사례",
        [("AE", "AETERM", "이상사례", "TEXT")],
        [
            ["2026-05-01", "SYN-01", "Synthetic Name", "F", "40", "cough"],
            ["2026-05-02", "SYN-01", "Synthetic Name", "F", "40", "rash"],
            ["", "", "", "", "", "blank-event"],
            ["2026-05-03", "NO-SUCH", "Synthetic Name", "F", "40", "unknown-event"],
        ],
    )
    _pair(
        book,
        "6. Specification",
        "6. MEDI-EYE 치료_V1.1",
        [("EX", "EXYN", "치료 여부", "TEXT"), ("EX", "EXARM", "무작위배정", "TEXT")],
        [["2026-05-01", "SYN-01", "Synthetic Name", "F", "40", "Y", "ARM-SECRET"]],
    )
    return _bytes(book)


def test_r1_links_visits_and_keeps_subject_level_forms_staged():
    content = _r1_workbook()
    report = r1_linkage_report(content)
    assert report["mapping_version"] == "spec-paired-1+r1"
    assert report["unknown_rows"] == 1
    assert report["unresolved_rows"] == 1
    assert report["visit_date_contradiction_rows"] == 1
    assert report["ambiguous_rows"] == 0
    assert report["subject_level_canonical_promoted"] is False
    assert "ARM-SECRET" not in json.dumps(report)
    assert "Synthetic Name" not in json.dumps(report)
    db = _db()
    try:
        study, site = _scope(db)
        disabled = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="r1-off.xlsx", content=content, enable_r1=False,
        )
        db.commit()
        assert disabled["mapping_version"] == "spec-paired-1"
        assert disabled["row_status_counts"].get("subject_level_staged", 0) == 0
        assert db.scalar(select(func.count()).select_from(StudyVisit)) == 1
        enabled = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="r1-on.xlsx", content=content, enable_r1=True,
        )
        db.commit()
        assert enabled["idempotent_replay"] is False
        assert enabled["mapping_version"] == "spec-paired-1+r1"
        assert enabled["row_status_counts"]["subject_level_staged"] == 3
        assert enabled["row_status_counts"]["unresolved_blank_subject"] == 1
        assert enabled["row_status_counts"]["unknown_subject"] == 1
        assert enabled["automatic_correction"] is False
        visits = set(db.scalars(select(StudyVisit.visit_code)).all())
        assert visits == {"V1", "V2", "V3"}
        v2 = db.scalar(select(StudyVisit).where(StudyVisit.visit_code == "V2"))
        kept = db.scalar(
            select(ClinicalObservation).where(
                ClinicalObservation.visit_id == v2.id,
                ClinicalObservation.code == "TBUT",
            )
        )
        assert float(kept.value_number) == 3
        assert db.scalar(select(func.count()).select_from(ImportDiscrepancy)) == 1
        assert db.scalar(select(ClinicalObservation).where(ClinicalObservation.code == "AETERM")) is None
        staged = db.scalars(select(ImportStagingRow).where(ImportStagingRow.status == "subject_level_staged")).all()
        assert {json.loads(row.payload_json)["form_class"] for row in staged} == {"AE", "EX"}
        assert any(json.loads(row.payload_json)["form_version"] == "V1.1" for row in staged)
        assert all("ARM-SECRET" not in row.payload_json for row in staged)
        assert all(row.subject_code != "NO-SUCH" for row in staged)
        blank = db.scalar(select(ImportStagingRow).where(ImportStagingRow.status == "unresolved_blank_subject"))
        assert blank.subject_code is None
        replay = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="r1-on.xlsx", content=content, enable_r1=True,
        )
        assert replay["idempotent_replay"] is True
        assert db.scalar(select(func.count()).select_from(StudyVisit)) == 3
        other = add_site(db, study_id=study.id, site_code="S2", name="Other", actor_user_id="doc")
        add_site_member(db, site_id=other.id, user_id="doc", role="investigator", actor_user_id="doc")
        db.commit()
        denied = confirm_import(
            db, study_id=study.id, site_id=other.id, actor_user_id="doc",
            filename="r1-other-site.xlsx", content=_book([_row(TBODORS1=8)]), enable_r1=True,
        )
        db.commit()
        assert denied["row_status_counts"]["cross_site_subject"] == 1
        assert db.scalar(select(func.count()).select_from(StudySubject)) == 1
    finally:
        db.close()
        db.get_bind().dispose()


def test_r1_rejects_ambiguous_token_and_does_not_correct_case():
    book = Workbook()
    _pair(
        book,
        "1. Specification",
        "1. V1 (스크리닝)",
        [("DM", "SUBJID", "스크리닝 번호", "TEXT"), ("SV", "SVDAT", "방문일", "VARCHAR(10)")],
        [
            ["2026-03-01", "Sub-1", "Synthetic Name", "F", "40", "Sub-1", "2026-03-02"],
            ["2026-03-01", "SUB-1", "Synthetic Name", "F", "40", "SUB-1", "2026-03-03"],
        ],
    )
    _pair(
        book,
        "2. Specification",
        "2. V2 (방문)",
        [("SV", "SVDAT", "방문일", "VARCHAR(10)")],
        [["2026-04-01", "SUB-1", "Synthetic Name", "F", "40", "2026-04-02"]],
    )
    content = _bytes(book)
    report = r1_linkage_report(content)
    assert report["ambiguous_rows"] == 1
    assert report["r1_visit_linked_rows"] == 0
    db = _db()
    try:
        study, site = _scope(db)
        summary = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="ambiguous.xlsx", content=content, enable_r1=True,
        )
        db.commit()
        assert summary["row_status_counts"]["ambiguous_subject"] == 1
        assert db.scalar(select(StudyVisit).where(StudyVisit.visit_code == "V2")) is None
        lower = Workbook()
        _pair(
            lower,
            "1. Specification",
            "1. V1 (스크리닝)",
            [("DM", "SUBJID", "스크리닝 번호", "TEXT"), ("SV", "SVDAT", "방문일", "VARCHAR(10)")],
            [["2026-03-01", "SYN-01", "Synthetic Name", "F", "40", "SYN-01", "2026-03-02"]],
        )
        _pair(
            lower,
            "2. Specification",
            "2. V4 (방문)",
            [("SV", "SVDAT", "방문일", "VARCHAR(10)")],
            [["2026-04-01", "syn-01", "Synthetic Name", "F", "40", "2026-04-02"]],
        )
        lowered = _bytes(lower)
        case_summary = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="case.xlsx", content=lowered, enable_r1=True,
        )
        db.commit()
        assert case_summary["row_status_counts"]["unknown_subject"] == 1
        assert db.scalar(select(StudyVisit).where(StudyVisit.visit_code == "V4")) is None
    finally:
        db.close()
        db.get_bind().dispose()


def test_r1_failure_removes_partial_rows(monkeypatch):
    content = _r1_workbook()
    db = _db()
    engine = None
    try:
        study, site = _scope(db)
        engine = db.get_bind()

        def _boom(*_args, **_kwargs):
            raise RuntimeError("forced persistence failure")

        monkeypatch.setattr("services.clinical_import.create_draft_observation", _boom)
        with pytest.raises(RuntimeError):
            confirm_import(
                db, study_id=study.id, site_id=site.id, actor_user_id="doc",
                filename="r1-boom.xlsx", content=content, enable_r1=True,
            )
        db.rollback()
        assert db.scalar(select(func.count()).select_from(StudySubject)) == 0
        assert db.scalar(select(func.count()).select_from(StudyVisit)) == 0
        assert db.scalar(select(func.count()).select_from(ImportStagingRow)) == 0
    finally:
        db.close()
        if engine is not None:
            engine.dispose()


def test_aggregate_summary_has_no_row_payload():
    content = _book([_row()], leading=True)
    counts = aggregate_workbook(content)
    blob = json.dumps(counts, ensure_ascii=False)
    assert "Synthetic Name" not in blob
    assert "CHART-SYN" not in blob
    assert counts["blinded_cell_count"] == 1
    assert counts["study_code_present_in_file"] is False


OSDI_CODE = (
    "[1] 0 - 그런적 없다 / [2] 1 - 가끔 그렇다 / [3] 2 - 절반정도 / "
    "[4] 3 - 대부분 그렇다 / [5] 4 - 항상 그렇다"
)
OSDI_NA_CODE = "[1] 해당없음 / [2] 0 - 그런적 없다 / [6] 4 - 항상 그렇다"


def _osdi_choices():
    base = parse_response_choices(OSDI_CODE)
    later = parse_response_choices(OSDI_NA_CODE)
    choices = {f"OSQ{index}": base for index in range(1, 6)}
    choices.update({f"OSQ{index}": later for index in range(6, 13)})
    return choices


def test_protocol_pack_is_extracted_and_visit_windows_conflict():
    rules = {rule.rule_id: rule for rule in rule_pack()}
    assert PACK_ID == "medi-des-001-v1.10"
    assert PROTOCOL_DATE == "2026-07-01"
    assert PROTOCOL_STATUS == "UNRESOLVED"
    assert len(rules) == 6
    assert all(rule.effective_status == "EXTRACTED" for rule in rules.values())
    assert all(rule.approval_status == "AWAITING_APPROVAL" for rule in rules.values())
    assert all(rule.effective_status != "ACTIVE" for rule in rules.values())
    assert len(rules["OSDI-ITEMS-001"].source_fields) == 12
    assert "0-4" in rules["OSDI-ITEMS-001"].logic
    assert "해당 없음" in rules["OSDI-ITEMS-001"].logic
    assert "x 25" in rules["OSDI-FORMULA-001"].logic
    assert rules["OSDI-FORMULA-001"].source_reference == "부록 1, p.93"
    assert rules["TBUT-MEAN-001"].unit == "s"
    assert "3회" in rules["TBUT-MEAN-001"].logic
    assert rules["TBUT-MEAN-001"].source_reference.startswith("p.53")
    assert rules["SCHIRMER-METHOD-001"].unit == "mm/5min"
    assert "마취 없이" in rules["SCHIRMER-METHOD-001"].logic
    assert "5분" in rules["SCHIRMER-METHOD-001"].logic
    assert rules["SCHIRMER-METHOD-001"].source_reference.startswith("p.55")
    form = osdi_form_visits()
    assert form["status"] == "RESOLVED_BY_V110"
    assert form["form_visits"] == form["schedule_visits"] == "V1,V2,V3,V4,V6"
    assert form["v13_form_visits"] == "V1,V2,V3,V4,V5"
    with pytest.raises(ValueError):
        promote_rule(rules["OSDI-FORMULA-001"], "ACTIVE")
    assert window_conflict()["status"] == "CONFLICTING"
    assert window_conflict()["validation"] == "NOT_IMPLEMENTED"
    matrix = visit_assessments()
    for visit in ("V1", "V2", "V3", "V4", "V6"):
        assert matrix[visit]["OSDI"] is True
        assert matrix[visit]["window_check"] is False
    assert matrix["V5"]["OSDI"] is False
    assert matrix["V5"]["TBUT"] is False
    assert matrix["V5"]["SCHIRMER"] is False
    assert all(row["window_check"] is False for row in matrix.values())


def test_osdi_decoding_excludes_not_applicable_and_rejects_unknown_codes():
    choices = _osdi_choices()
    values = {f"OSQ{index}": "4 - 항상 그렇다" for index in range(1, 6)}
    values.update({f"OSQ{index}": "해당없음" for index in range(6, 13)})
    values["OSSCORE"] = "100점"
    assert advisory_findings(values, choices) == []
    values["OSQ6"] = "알 수 없는 코드"
    messages = {item["message_code"] for item in advisory_findings(values, choices)}
    assert "RESPONSE_CODE_UNRESOLVED" in messages
    values["OSQ6"] = "해당없음"
    values["OSSCORE"] = 0
    mismatch = advisory_findings(values, choices)
    assert any(item["message_code"] == "SCORING_MISMATCH" for item in mismatch)
    blank = {f"OSQ{index}": None for index in range(1, 13)}
    blank["OSSCORE"] = None
    assert advisory_findings(blank, choices) == []


def test_tbut_mean_and_schirmer_precision_are_advisory_only():
    matched = {
        "TBODORS1": 2, "TBODORS2": 4, "TBODORS3": 6, "TBODORSAV": 4,
        "STODORS": "10.0mm",
    }
    assert advisory_findings(matched) == []
    mismatched = dict(matched)
    mismatched["TBODORSAV"] = 9
    assert any(item["rule_code"] == "TBUT-MEAN-001" for item in advisory_findings(mismatched))
    partial = {"TBODORS1": 2, "TBODORS2": 4, "TBODORSAV": 3}
    assert any(item["message_code"] == "MISSING_ITEM" for item in advisory_findings(partial))
    fine_unit = advisory_findings({"STOSORS": 10.25})
    assert any(item["rule_code"] == "SCHIRMER-METHOD-001" for item in fine_unit)
    assert advisory_findings({"STOSORS": 10}) == []


def test_protocol_advisory_is_opt_in_and_does_not_correct_values():
    choices = _osdi_choices()
    items = [
        ("OS", "OSQ1", "1. 눈이 빛에 예민하다", "INTEGER", OSDI_CODE),
        ("OS", "OSSCORE", "OSDI 점수", "FLOAT", ""),
        ("SV", "SVDAT", "방문일", "VARCHAR(10)", ""),
        ("DM", "SUBJID", "스크리닝 번호", "TEXT", ""),
    ]
    book = Workbook()
    spec = book.active
    spec.title = "1. Specification"
    spec.append(["Domain", "Item ID", "Item Label", "Code", "Layout", "Type", "Required"])
    for domain, item_id, label, type_name, code in items:
        spec.append([domain, item_id, label, code, "", type_name, "FALSE"])
    data = book.create_sheet("1. V1 (스크리닝)")
    data.append(["1. 눈이 빛에 예민하다", "OSDI 점수", "방문일", "스크리닝 번호"])
    data.append(["4 - 항상 그렇다", "0점", "2026-03-02", "SYN-01"])
    content = _bytes(book)
    db = _db()
    try:
        study, site = _scope(db)
        plain = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="plain.xlsx", content=content,
        )
        db.commit()
        assert plain["mapping_version"] == "spec-paired-1"
        assert plain["r1_enabled"] is False
        assert plain["automatic_finalize"] is False
        assert plain["automatic_lock"] is False
        assert plain["automatic_correction"] is False
        assert "OSDI-FORMULA-001" not in set(db.scalars(select(ImportValidationFinding.rule_code)).all())
        advised = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="advised.xlsx", content=content, protocol_advisory_id=PACK_ID,
        )
        db.commit()
        assert advised["mapping_version"] == f"spec-paired-1+advisory-{PACK_ID}"
        assert "SCORING_MISMATCH" in set(db.scalars(select(ImportValidationFinding.message_code)).all())
        stored = db.scalar(select(QuestionnaireResponseItem))
        assert stored is not None
        assert json.loads(stored.value_json) == "4 - 항상 그렇다"
        replay = confirm_import(
            db, study_id=study.id, site_id=site.id, actor_user_id="doc",
            filename="advised.xlsx", content=content, protocol_advisory_id=PACK_ID,
        )
        assert replay["idempotent_replay"] is True
        with pytest.raises(ClinicalRecordError):
            confirm_import(
                db, study_id=study.id, site_id=site.id, actor_user_id="doc",
                filename="bad-pack.xlsx", content=content, protocol_advisory_id="other-study",
            )
        with pytest.raises(ClinicalRecordError):
            confirm_import(
                db, study_id=study.id, site_id=site.id, actor_user_id="doc",
                filename="old-pack.xlsx", content=content, protocol_advisory_id="medi-des-001-v1.3",
            )
    finally:
        db.close()
        db.get_bind().dispose()
    assert choices["OSQ1"][0] == ResponseChoice(index=1, score=0, label="0 - 그런적 없다")

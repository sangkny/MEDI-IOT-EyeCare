"""명세 시트와 데이터 시트가 짝인 표 형식 반출을 읽기 전용으로 해석한다.

시트 이름이나 특정 EDC 제품 이름을 고정하지 않는다.
열 순서가 명세 라벨과 일치할 때만 매핑을 확정한다.
"""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook

MAPPING_VERSION = "spec-paired-1"
ADAPTER_NAME = "specification_paired"
LEADING_HEADERS = ("작성일자", "병록번호", "이름", "성별", "나이")
SPEC_HEADERS = ("Domain", "Item ID", "Item Label", "Code", "Layout", "Type", "Required")
SENSITIVE_LABELS = {"이름", "이니셜"}
BLINDED_DOMAINS = {"RN"}
VISIT_CODE_RE = re.compile(r"V(\d+)\s*\(")


@dataclass
class ItemSpec:
    domain: str
    item_id: str
    label: str
    type_name: str
    sensitive: bool
    blinded: bool


@dataclass
class ParsedRow:
    sheet_name: str
    sheet_index: int
    source_row: int
    visit_code: str | None
    values: dict[str, str | float | None]
    sensitive_hashes: dict[str, str]
    blinded_item_ids: list[str] = field(default_factory=list)


@dataclass
class ParsedSheet:
    name: str
    index: int
    alignment: str
    visit_code: str | None
    items: list[ItemSpec]
    rows: list[ParsedRow]
    mismatch_count: int = 0


@dataclass
class ParsedWorkbook:
    content_sha256: str
    byte_size: int
    sheets: list[ParsedSheet]

    @property
    def data_row_count(self) -> int:
        return sum(len(sheet.rows) for sheet in self.sheets)


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _missing(value: object) -> bool:
    return _text(value) == ""


def _sensitive_label(label: str) -> bool:
    if label in SENSITIVE_LABELS:
        return True
    return "생년월일" in label


def _blinded(domain: str, label: str) -> bool:
    if domain in BLINDED_DOMAINS:
        return True
    return "무작위" in label or "배정" in label


def _hash_value(value: object) -> str:
    return hashlib.sha256(_text(value).encode("utf-8")).hexdigest()


def _public_value(value: object) -> str | float | None:
    if _missing(value):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return float(value)
    return _text(value)


def _spec_items(rows: list[tuple]) -> list[ItemSpec] | None:
    if not rows:
        return None
    header = tuple(_text(cell) for cell in rows[0][: len(SPEC_HEADERS)])
    if header != SPEC_HEADERS:
        return None
    items: list[ItemSpec] = []
    for row in rows[1:]:
        if row is None or _missing(row[0] if row else None):
            continue
        label = _text(row[2] if len(row) > 2 else "")
        domain = _text(row[0])
        items.append(ItemSpec(
            domain=domain,
            item_id=_text(row[1] if len(row) > 1 else ""),
            label=label,
            type_name=_text(row[5] if len(row) > 5 else ""),
            sensitive=_sensitive_label(label),
            blinded=_blinded(domain, label),
        ))
    return items


def _alignment_offset(headers: list[str], items: list[ItemSpec]) -> int | None:
    labels = [item.label for item in items]
    if not labels:
        return None
    best: tuple[int, int] | None = None
    for offset in range(0, 9):
        mismatches = 0
        for index, label in enumerate(labels):
            position = offset + index
            if position >= len(headers) or headers[position] != label:
                mismatches += 1
        if best is None or mismatches < best[0]:
            best = (mismatches, offset)
    if best is None or best[0] != 0:
        return None
    return best[1]


def _visit_code(sheet_name: str) -> str | None:
    match = VISIT_CODE_RE.search(sheet_name)
    if match is None:
        return None
    return f"V{match.group(1)}"


def _read_sheet_rows(worksheet) -> list[tuple]:
    return [tuple(row) for row in worksheet.iter_rows(values_only=True)]


def parse_workbook(content: bytes) -> ParsedWorkbook:
    if not content:
        raise ValueError("빈 파일은 해석할 수 없습니다.")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        names = list(workbook.sheetnames)
        cached = {name: _read_sheet_rows(workbook[name]) for name in names}
    finally:
        workbook.close()
    sheets: list[ParsedSheet] = []
    index = 0
    while index < len(names):
        items = _spec_items(cached[names[index]])
        if items is not None and index + 1 < len(names):
            data_name = names[index + 1]
            data_rows = cached[data_name]
            headers = [_text(cell) for cell in data_rows[0]] if data_rows else []
            offset = _alignment_offset(headers, items)
            visit_code = _visit_code(data_name)
            parsed_rows: list[ParsedRow] = []
            if offset is not None:
                for row_number, row in enumerate(data_rows[1:], start=2):
                    if row is None or all(_missing(cell) for cell in row):
                        continue
                    values: dict[str, str | float | None] = {}
                    hashes: dict[str, str] = {}
                    blinded_ids: list[str] = []
                    for item_index, item in enumerate(items):
                        position = offset + item_index
                        cell = row[position] if position < len(row) else None
                        if item.sensitive:
                            hashes[item.item_id] = _hash_value(cell)
                            continue
                        if item.blinded:
                            blinded_ids.append(item.item_id)
                            hashes[item.item_id] = _hash_value(cell)
                            continue
                        values[item.item_id] = _public_value(cell)
                    if offset == len(LEADING_HEADERS) and tuple(headers[:offset]) == LEADING_HEADERS:
                        for lead_index, lead_name in enumerate(LEADING_HEADERS):
                            if lead_name in {"병록번호", "이름"}:
                                hashes[lead_name] = _hash_value(row[lead_index] if lead_index < len(row) else None)
                    parsed_rows.append(ParsedRow(
                        sheet_name=data_name,
                        sheet_index=index + 1,
                        source_row=row_number,
                        visit_code=visit_code,
                        values=values,
                        sensitive_hashes=hashes,
                        blinded_item_ids=blinded_ids,
                    ))
                sheets.append(ParsedSheet(
                    name=data_name,
                    index=index + 1,
                    alignment="confirmed",
                    visit_code=visit_code,
                    items=items,
                    rows=parsed_rows,
                ))
            else:
                sheets.append(ParsedSheet(
                    name=data_name,
                    index=index + 1,
                    alignment="unresolved",
                    visit_code=visit_code,
                    items=items,
                    rows=[],
                    mismatch_count=-1,
                ))
            index += 2
            continue
        index += 1
    return ParsedWorkbook(
        content_sha256=hashlib.sha256(content).hexdigest(),
        byte_size=len(content),
        sheets=sheets,
    )
